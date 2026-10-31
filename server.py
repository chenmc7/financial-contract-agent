# -*- coding: utf-8 -*-
"""ContractGuard 企业级金融合同全生命周期风险管理服务。"""
import json
import os
import re
import time
import uuid
import csv
import io
import hmac
import tempfile
from urllib.parse import urlparse

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import Optional, List, Dict

from engines import (parser, costlens, scenario, covenant, diffengine, ai_client,
                     samples, assessment, orchestrator, riskengine, file_parser)
from engines.uparse import _attach_evidence_spans
from engines.uparse import transactions as transaction_indexer
from storage import Store

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.abspath(os.path.expanduser(os.getenv(
    "CONTRACTGUARD_UPLOAD_DIR", os.path.join(BASE_DIR, "uploads"))))
os.makedirs(UPLOAD_DIR, exist_ok=True)
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
DATA_DIR = os.path.join(BASE_DIR, "data")
DB_PATH = os.path.abspath(os.path.expanduser(os.getenv(
    "CONTRACTGUARD_DB_PATH", os.path.join(DATA_DIR, "contractguard.db"))))
os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
DEMO_MODE = os.getenv("CONTRACTGUARD_DEMO_MODE", "0").lower() in ("1", "true", "yes")
SHARE_MODE = os.getenv("CONTRACTGUARD_SHARE_MODE", "0").lower() in ("1", "true", "yes")

# ---------- 公网只读保护（SHARE_MODE）----------
# 公网隧道实例与本地实例共用同一份数据库，任何访客的写操作都会落到真实数据上。
# 因此 SHARE_MODE 下只放行只读方法，写方法一律拦截，只有下面这份"纯计算接口"
# 白名单例外——它们不落库、不改配置，是演示必须的能力。
SHARE_WRITE_METHODS = {"POST", "PUT", "PATCH", "DELETE"}
SHARE_READONLY_POST_PATHS = {
    "/api/analyze/costlens",       # 只读测算
    "/api/analyze/scenario",       # 只读情景模拟
    "/api/analyze/covenant",       # 只读契约判定（share 模式下不持久化财务数据）
    "/api/analyze/diff",           # 只读版本比对
    "/api/analyze/comprehensive",  # 只读一键评估
    "/api/agent/tasks",            # 只读 Agent 编排（仅追加审计日志）
    "/api/ai/chat",                # 只读问答（不落库）
    "/api/ai/chat-once",
}
# 可选访问令牌：配置后，请求头携带正确令牌即可临时放行写操作（维护者远程使用）。
SHARE_TOKEN_ENV = "CONTRACTGUARD_SHARE_TOKEN"
SHARE_TOKEN_HEADER = "X-Share-Token"
SHARE_TOKEN = (os.getenv(SHARE_TOKEN_ENV) or "").strip()
SHARE_READONLY_MESSAGE = ("公网演示实例为只读模式：删除、改名、上传、修改配置等写操作已被拦截。"
                          "请在本地实例（127.0.0.1）执行写操作后再刷新公网页面。")

app = FastAPI(title="ContractGuard API", version="1.0.0")
ai_client.load_config(CONFIG_PATH)
STORE = Store(DB_PATH)
SECURITY = {"deterministic_only": SHARE_MODE, "external_scope": "masked_evidence_only"}

def _load_contract_cache():
    """从 SQLite 重载合同缓存，让本机与公网进程共享最新数据。"""
    loaded = STORE.load_contracts()
    if not DEMO_MODE:
        # Older versions persisted built-in demo contracts unconditionally. Preserve
        # those database rows, but do not expose or analyze them in real-data mode.
        loaded = {cid: rec for cid, rec in loaded.items() if rec.get("source") != "builtin"}
    return loaded


# 合同仓库：id -> {id, name, type, source, text, parsed}
CONTRACTS = _load_contract_cache()


def _refresh_persisted_evidence():
    """升级旧记录中的证据坐标，避免服务重启后继续展示历史截断片段。"""
    for rec in CONTRACTS.values():
        parsed = rec.get("parsed") or {}
        fields = parsed.get("fields") or {}
        if not fields or not rec.get("text"):
            continue
        try:
            parsed["document"] = _attach_evidence_spans(fields, rec["text"])
            rebuilt_transactions = transaction_indexer.build_transactions(
                rec["text"], fields, parsed.get("contract_type") or {})
            parsed["transactions"] = rebuilt_transactions.get("transactions", [])
            parsed["transaction_count"] = rebuilt_transactions.get("count", 0)
            auth = fields.get("info_auth") or {}
            auth_value = auth.get("value") if isinstance(auth.get("value"), dict) else {}
            retention = auth_value.get("retention_years")
            # 旧版本曾从整个公告段落抓取“2025 年”作为保留期限；仅对未人工确认的
            # 可疑字段用最新规则重算，保留其他用户修正和分析结果。
            if auth.get("status") != "confirmed_by_user" and isinstance(retention, (int, float)) and (retention > 120 or 1900 <= retention <= 2100):
                refreshed = parser.parse_contract(rec["text"], rec.get("name") or "未命名合同", use_llm=False)
                fresh_auth = (refreshed.get("fields") or {}).get("info_auth")
                if fresh_auth:
                    fields["info_auth"] = fresh_auth
            rec["parsed"] = parsed
            STORE.save_contract(rec)
        except Exception:
            # 单条旧记录升级失败不影响合同库启动；新上传仍使用最新绑定逻辑。
            continue


_refresh_persisted_evidence()


@app.middleware("http")
async def refresh_contract_cache(request, call_next):
    """每次请求前读取共享数据库，避免公网进程保留旧的内存合同列表。"""
    global CONTRACTS
    CONTRACTS = _load_contract_cache()
    return await call_next(request)


def share_guard_decision(method, path, token=None):
    """公网只读保护判定：返回 None 表示放行，否则返回 403 响应体。

    抽成纯函数，便于单元测试逐条覆盖判定规则（不依赖 ASGI 请求对象）。
    """
    if not SHARE_MODE:
        return None
    method = (method or "").upper()
    if method not in SHARE_WRITE_METHODS:
        return None
    # 归一化尾部斜杠后再比对白名单，避免 /api/analyze/costlens/ 绕过
    normalized = "/" + (path or "").strip("/")
    if method == "POST" and normalized in SHARE_READONLY_POST_PATHS:
        return None
    if SHARE_TOKEN and hmac.compare_digest(token or "", SHARE_TOKEN):
        return None
    return {
        "ok": False, "share_mode": True, "readonly": True,
        "method": method, "path": path,
        "error": SHARE_READONLY_MESSAGE,
        "hint": "维护者可通过环境变量 %s 配置访问令牌，并在请求头 %s 中携带后放行写操作。"
                % (SHARE_TOKEN_ENV, SHARE_TOKEN_HEADER),
    }


@app.middleware("http")
async def share_readonly_guard(request, call_next):
    """公网只读保护：SHARE_MODE 下拦截所有写请求。

    判定顺序：
    1. 非 SHARE_MODE（本地开发实例）→ 全部放行，不影响任何人改代码；
    2. GET/HEAD/OPTIONS 等只读方法 → 放行；
    3. 白名单里的纯计算接口（POST）→ 放行，保证公网仍能演示测算；
    4. 携带正确 X-Share-Token 的请求 → 放行（配置了令牌才生效）；
    5. 其余 POST/PUT/PATCH/DELETE → 403 + 可读提示。

    该中间件注册在缓存刷新之后，因此是最外层，被拦截的写请求连数据库都不会读。
    """
    decision = share_guard_decision(request.method, request.url.path,
                                    request.headers.get(SHARE_TOKEN_HEADER))
    if decision is None:
        return await call_next(request)
    return JSONResponse(status_code=403, content=decision)


# ---------- 工具 ----------

def _add_contract(cid, name, ctype, source, text, meta=None, use_llm=False, file_diagnostics=None,
                  file_path=None, original_filename=None, display_name=None):
    parsed = parser.parse_contract(text, name, use_llm=use_llm)
    if meta:
        # 样例合同：用人工核对的元数据补强解析结果
        if meta.get("principal") and parsed.get("principal") is None:
            parsed["principal"] = meta["principal"]
        if meta.get("annual_rate") is not None and parsed.get("annual_rate") is None:
            parsed["annual_rate"] = meta["annual_rate"]
        if meta.get("term_months") and parsed.get("term_months") is None:
            parsed["term_months"] = meta["term_months"]
        if meta.get("repayment"):
            parsed["repayment"] = meta["repayment"]
        if meta.get("fees"):
            parsed["fees"] = meta["fees"]
        if meta.get("prepayment_penalty") and not parsed.get("prepayment_penalty"):
            p = meta.get("principal") or parsed.get("principal") or 0
            parsed["prepayment_penalty"] = {"rate": meta["prepayment_penalty"],
                                           "amount": p * meta["prepayment_penalty"],
                                           "source": "合同约定（样例标注）"}
        if meta.get("guarantee") and parsed.get("guarantee") == "未明确约定":
            parsed["guarantee"] = meta["guarantee"]
        parsed["confidence"] = 1.0
    display_name = (display_name or name or "未命名合同").strip()[:80]
    rec = {"id": cid, "name": display_name, "display_name": display_name,
           "original_filename": original_filename or name, "type": ctype, "source": source,
           "text": text, "parsed": parsed, "created_at": time.time(),
           "file_diagnostics": file_diagnostics or {"format": "text", "status": "success"},
           "file_path": file_path, "access_level": "owner", "analysis_version": 1}
    CONTRACTS[cid] = rec
    STORE.save_contract(rec)
    STORE.audit("contract_created", cid, {"source": source, "type": ctype})
    return rec


def _public(rec, detail=False):
    parsed = rec["parsed"]
    ctype = parsed.get("contract_type") or {}
    fields = parsed.get("fields") or {}
    out = {
        "id": rec["id"], "name": rec.get("display_name") or rec["name"],
        "display_name": rec.get("display_name") or rec["name"],
        "original_filename": rec.get("original_filename") or rec["name"],
        "type": rec["type"], "source": rec["source"],
        "principal": parsed.get("principal"), "annual_rate": parsed.get("annual_rate"),
        "term_months": parsed.get("term_months"), "repayment": parsed.get("repayment"),
        "guarantee": parsed.get("guarantee"), "fees": parsed.get("fees", []),
        "prepayment_penalty": parsed.get("prepayment_penalty"),
        "covenants": parsed.get("covenants", []),
        "confidence": parsed.get("confidence"), "word_count": parsed.get("word_count"),
        "clause_count": len(parsed.get("clauses", [])),
        "created_at": rec["created_at"],
        "analysis_version": rec.get("analysis_version", 1),
        "access_level": rec.get("access_level", "owner"),
        "file_diagnostics": rec.get("file_diagnostics") or {},
        # 通用结构化解析结果
        "type_key": ctype.get("key"),
        "type_label": ctype.get("label") or rec["type"],
        "type_confidence": ctype.get("confidence"),
        "type_evidence": ctype.get("evidence"),
        "document_type": parsed.get("document_type"),
        "transaction_type": parsed.get("transaction_type"),
        "transaction_count": parsed.get("transaction_count", 0),
        "field_count": parsed.get("field_count"),
        "pipeline": parsed.get("pipeline"),
        "field_status": {key: (fields.get(key) or {}).get("status")
                         for key in ("amount", "currency", "rate", "term", "payment_method")},
    }
    if detail:
        out["fields"] = fields
        out["document"] = parsed.get("document") or {}
        out["transactions"] = parsed.get("transactions") or []
    return out


# ---------- 初始化样例合同 ----------

if DEMO_MODE:
    for sid, sc in samples.SAMPLE_CONTRACTS.items():
        if sid not in CONTRACTS:
            _add_contract(sid, sc["name"], sc.get("type", "合同"), "builtin", sc["content"], sc.get("meta"))


# ---------- API：合同管理 ----------

class ImportTextBody(BaseModel):
    name: str
    text: str


class RenameBody(BaseModel):
    display_name: str


@app.get("/api/contracts")
def list_contracts():
    return {"contracts": [_public(r) for r in CONTRACTS.values()]}


@app.patch("/api/contracts/{cid}/name")
def rename_contract(cid: str, body: RenameBody):
    rec = CONTRACTS.get(cid)
    if not rec:
        raise HTTPException(404, "合同不存在")
    new_name = (body.display_name or "").strip()
    if not new_name:
        raise HTTPException(400, "名称不能为空")
    if len(new_name) > 80:
        raise HTTPException(400, "名称不能超过 80 个字符")
    old_name = rec.get("display_name") or rec.get("name") or "未命名合同"
    if new_name == old_name:
        return {"ok": True, "contract": _public(rec)}
    # 记录轻量重命名历史；原始文件名和合同 ID 永不改变。
    STORE.revision(cid, "display_name", old_name, new_name, "user_rename")
    rec["display_name"] = new_name
    rec["name"] = new_name  # 兼容所有仍消费 name 的页面和分析接口
    rec["analysis_version"] = rec.get("analysis_version", 1) + 1
    STORE.save_contract(rec)
    STORE.audit("contract_renamed", cid, {"old_display_name": old_name, "display_name": new_name})
    return {"ok": True, "contract": _public(rec)}


@app.get("/api/contracts/{cid}")
def get_contract(cid: str):
    rec = CONTRACTS.get(cid)
    if not rec:
        raise HTTPException(404, "合同不存在")
    try:
        orch = orchestrator.orchestrate(rec["parsed"])
    except Exception as exc:  # 调度失败不阻塞详情展示
        orch = {"plan": {"steps": [], "cost_ready": False,
                         "advice": "Agent 调度异常：%s" % exc},
                "review": [], "risks": [], "risk_summary": {},
                "diagnostics": rec["parsed"].get("diagnostics") or {},
                "consistency": rec["parsed"].get("consistency") or {"conflicts": []}}
    return {"contract": _public(rec, detail=True), "text": rec["text"],
            "clauses": rec["parsed"].get("clauses", []), "orchestration": orch,
            "revisions": STORE.revisions(cid), "review_status": STORE.review_statuses(cid)}


@app.delete("/api/contracts/{cid}")
def delete_contract(cid: str):
    if cid in samples.SAMPLE_CONTRACTS:
        raise HTTPException(400, "内置样例合同不可删除")
    if cid not in CONTRACTS:
        raise HTTPException(404, "合同不存在")
    rec = CONTRACTS.pop(cid)
    if rec.get("file_path") and os.path.exists(rec["file_path"]):
        os.unlink(rec["file_path"])
    STORE.delete_contract(cid)
    return {"ok": True, "purged": True}


@app.post("/api/contracts/upload")
async def upload_contract(file: UploadFile = File(...)):
    data = await file.read()
    if not data:
        raise HTTPException(400, "上传失败：文件内容为空")
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(413, "上传失败：文件超过 15MB 限制")
    try:
        extracted = file_parser.extract_document(file.filename, data, file.content_type)
    except PermissionError as exc:
        raise HTTPException(422, "解析失败：%s" % exc)
    except ValueError as exc:
        raise HTTPException(415, "上传失败：%s" % exc)
    except RuntimeError as exc:
        raise HTTPException(503, "解析服务不可用：%s" % exc)
    except Exception as exc:
        raise HTTPException(422, "解析失败：%s" % str(exc)[:160])
    text = extracted["text"]
    if not text or len(re.sub(r"\s", "", text)) < 30:
        status = (extracted.get("diagnostics") or {}).get("status")
        if status == "needs_text_version":
            raise HTTPException(422, "解析失败：PDF 为扫描版或无可提取文本，请上传文本型 PDF 或 DOCX")
        raise HTTPException(422, "解析失败：未提取到足够的合同文本")
    cid = "UP" + uuid.uuid4().hex[:8].upper()
    name = os.path.splitext(file.filename)[0][:60]
    ext = os.path.splitext(file.filename)[1].lower()
    path = os.path.join(UPLOAD_DIR, cid + ext)
    fd, tmp_path = tempfile.mkstemp(prefix=cid + "-", suffix=ext, dir=UPLOAD_DIR)
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
        os.replace(tmp_path, path)
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
    allow_llm = ai_client.is_configured() and not SECURITY["deterministic_only"]
    rec = _add_contract(cid, name, "上传合同", "upload", text, use_llm=allow_llm,
                        file_diagnostics=extracted["diagnostics"], file_path=path,
                        original_filename=file.filename, display_name=name)
    try:
        orch = orchestrator.orchestrate(rec["parsed"])
    except Exception:
        orch = None
    out = {"contract": _public(rec, detail=True), "clauses": rec["parsed"].get("clauses", [])}
    if orch:
        out["orchestration"] = orch
    return out


@app.post("/api/contracts/import-text")
def import_text(body: ImportTextBody):
    if len(re.sub(r"\s", "", body.text)) < 30:
        raise HTTPException(400, "文本过短，无法识别合同要素")
    cid = "TX" + uuid.uuid4().hex[:8].upper()
    allow_llm = ai_client.is_configured() and not SECURITY["deterministic_only"]
    rec = _add_contract(cid, body.name[:60] or "粘贴合同", "粘贴文本", "text", body.text, use_llm=allow_llm)
    try:
        orch = orchestrator.orchestrate(rec["parsed"])
    except Exception:
        orch = None
    out = {"contract": _public(rec, detail=True), "clauses": rec["parsed"].get("clauses", [])}
    if orch:
        out["orchestration"] = orch
    return out


class FieldUpdateBody(BaseModel):
    value: object = None
    display: Optional[str] = None
    source: str = "user_edit"
    resolution: str = "adopt_value"  # adopt_value / uncertain / legal_review


def _sync_legacy(parsed, key, value):
    if key == "amount":
        parsed["principal"] = float(value) if value not in (None, "") else None
    elif key == "actual_financing_amount":
        parsed["principal"] = float(value) if value not in (None, "") else None
        # 保理场景的通用金额字段与专属实际融资金额必须保持同一业务角色。
        if parsed.get("fields", {}).get("amount"):
            parsed["fields"]["amount"].update(value=parsed["principal"], found=value not in (None, ""),
                                                role="actual_financing_amount")
    elif key == "rate":
        v = float(value) if value not in (None, "") else None
        parsed["annual_rate"] = (v / 100.0 if v is not None and v >= 1 else v)
    elif key == "term": parsed["term_months"] = int(float(value)) if value not in (None, "") else None
    elif key == "single_financing_term":
        if value not in (None, "") and re.fullmatch(r"\d+(?:\.\d+)?", str(value).strip()):
            parsed["term_months"] = int(float(value))
            if parsed.get("fields", {}).get("term"):
                parsed["fields"]["term"].update(value=parsed["term_months"], found=True,
                                                  status="confirmed_by_user", role="actual_financing_term")
    elif key == "payment_method": parsed["repayment"] = str(value) if value not in (None, "") else None
    elif key == "guarantee": parsed["guarantee"] = str(value) if value not in (None, "") else None


def _contract_kind(rec):
    """合同的类型标识：优先解析结果里的类型 key，退回上传时记录的类型。"""
    return ((rec.get("parsed") or {}).get("contract_type") or {}).get("key") or rec.get("type") or ""


def _is_cost_comparable(rec):
    """能否参与情景对比：本金、利率、期限三项齐备。"""
    parsed = rec.get("parsed") or {}
    return bool(parsed.get("principal")) and parsed.get("annual_rate") is not None and bool(parsed.get("term_months"))


def _select_compare_peer(rec, contracts=None):
    """复核重算时的情景对比对手：同类型且创建时间最近的合同。

    返回 (合同记录, 是否同类型)；没有可用对手时返回 None。

    旧实现取 `CONTRACTS.values()` 里第一个可测算的合同，而 CONTRACTS 是从 SQLite
    重载的 dict，载入顺序一变对手就变，同一份合同两次重算的对比结论会不一致。
    这里显式按（类型是否相同, 创建时间倒序, 合同 ID 升序）排序，结果可复现：
    - 同类型优先，找不到同类型才退回其他类型的最近一份；
    - 创建时间相同（批量导入）时用 ID 兜底，避免依赖字典顺序。
    """
    pool = CONTRACTS if contracts is None else contracts
    own_id = rec.get("id")
    own_kind = _contract_kind(rec)
    candidates = []
    for other in pool.values():
        if other.get("id") == own_id or not _is_cost_comparable(other):
            continue
        same_type = bool(own_kind) and _contract_kind(other) == own_kind
        candidates.append((0 if same_type else 1, -(other.get("created_at") or 0), other.get("id") or "", same_type, other))
    if not candidates:
        return None
    candidates.sort(key=lambda item: (item[0], item[1], item[2]))
    _, _, _, same_type, peer = candidates[0]
    return peer, same_type


def _rerun_single(rec):
    orch = orchestrator.orchestrate(rec["parsed"])
    result = {"orchestration": orch}
    if "costlens" in orch.get("tool_results", {}):
        result["costlens"] = orch["tool_results"]["costlens"]
    result["risk"] = orch.get("tool_results", {}).get("riskengine")
    if _is_cost_comparable(rec):
        selected = _select_compare_peer(rec)
        if selected:
            peer, same_type = selected
            comparison = scenario.compare_contracts([rec, peer], "hold")
            for item in comparison:
                item["is_self"] = item.get("id") == rec["id"]
            result["scenario_compare"] = comparison
            # 对手显式标注，避免用户对着两个数字猜"跟谁比"
            result["scenario_compare_peer"] = {
                "id": peer["id"],
                "name": peer.get("display_name") or peer.get("name"),
                "type_key": _contract_kind(peer),
                "type_label": ((peer.get("parsed") or {}).get("contract_type") or {}).get("label") or peer.get("type"),
                "created_at": peer.get("created_at"),
                "same_type": same_type,
                "rule": "同类型且创建时间最近" if same_type else "无同类型可测算合同，退回创建时间最近的合同",
            }
    rec["analysis_version"] = rec.get("analysis_version", 1) + 1
    STORE.save_contract(rec)
    STORE.save_analysis(rec["id"], "auto_recompute", result, rec["analysis_version"])
    return result


@app.patch("/api/contracts/{cid}/fields/{field_key}")
def update_field(cid: str, field_key: str, body: FieldUpdateBody):
    rec = CONTRACTS.get(cid)
    if not rec: raise HTTPException(404, "合同不存在")
    fields = rec["parsed"].get("fields") or {}
    if field_key not in fields: raise HTTPException(404, "字段不存在")
    old = json.loads(json.dumps(fields[field_key], ensure_ascii=False))
    f = fields[field_key]
    history = f.setdefault("correction_history", [])
    normalized_value = body.value
    try:
        if field_key in {"amount", "actual_financing_amount", "credit_limit", "term", "credit_term"} and body.value not in (None, ""):
            normalized_value = float(body.value)
            if field_key in {"term", "credit_term"}:
                normalized_value = int(normalized_value)
        if field_key == "rate" and body.value not in (None, ""):
            raw_rate = float(body.value)
            normalized_value = raw_rate / 100.0 if raw_rate >= 1 else raw_rate
    except (TypeError, ValueError):
        raise HTTPException(400, "字段修正失败：请输入有效的数字")
    change = {"old_value": f.get("value"), "new_value": normalized_value, "source": body.source,
              "changed_at": time.time(), "original_evidence_preserved": True}
    history.append(change)
    if body.resolution in {"uncertain", "legal_review"}:
        f["status"] = "uncertain"
        f["note"] = "用户暂不确定，已进入人工复核队列" if body.resolution == "uncertain" else "已提交人工复核，暂不采信该字段"
    else:
        f.update({"found": True, "value": normalized_value, "display": body.display or str(normalized_value),
                  "status": "confirmed_by_user", "method": "user_correction", "confidence": 1.0})
        _sync_legacy(rec["parsed"], field_key, normalized_value)
    STORE.revision(cid, field_key, old, f, body.source)
    STORE.audit("field_updated", cid, {"field": field_key, "resolution": body.resolution})
    recomputed = _rerun_single(rec)
    return {"ok": True, "field": f, "revision": change, "recomputed": recomputed,
            "analysis_version": rec["analysis_version"]}


class ReviewBody(BaseModel):
    status: str
    note: str = ""


@app.post("/api/contracts/{cid}/reviews/{item_id}")
def update_review(cid: str, item_id: str, body: ReviewBody):
    if cid not in CONTRACTS: raise HTTPException(404, "合同不存在")
    allowed = {"submitted_legal", "confirmed", "accepted_risk", "pending", "material_change", "wording_only"}
    if body.status not in allowed: raise HTTPException(400, "无效复核状态")
    STORE.set_review(cid, item_id, body.status, body.note)
    STORE.audit("review_status_updated", cid, {"item": item_id, "status": body.status})
    return {"ok": True, "item_id": item_id, "status": body.status}


@app.post("/api/contracts/{cid}/reparse")
def reparse_contract(cid: str):
    rec = CONTRACTS.get(cid)
    if not rec: raise HTTPException(404, "合同不存在")
    corrections = {k: json.loads(json.dumps(v)) for k, v in (rec["parsed"].get("fields") or {}).items()
                   if v.get("status") == "confirmed_by_user"}
    parsed = parser.parse_contract(rec["text"], rec["name"], use_llm=ai_client.is_configured() and not SECURITY["deterministic_only"])
    for key, corrected in corrections.items():
        if key in parsed.get("fields", {}):
            parsed["fields"][key] = corrected
            _sync_legacy(parsed, key, corrected.get("value"))
    rec["parsed"] = parsed
    result = _rerun_single(rec)
    STORE.audit("contract_reparsed", cid, {"preserved_user_corrections": len(corrections)})
    return {"contract": _public(rec, detail=True), "recomputed": result}


@app.get("/api/contracts/{cid}/export")
def export_contract(cid: str):
    rec = CONTRACTS.get(cid)
    if not rec: raise HTTPException(404, "合同不存在")
    return JSONResponse({"contract": _public(rec, detail=True), "text": rec["text"],
                         "revisions": STORE.revisions(cid), "reviews": STORE.review_statuses(cid),
                         "financial_history": STORE.financial_history(cid)})


class AccessBody(BaseModel):
    access_level: str


@app.patch("/api/contracts/{cid}/access")
def update_access(cid: str, body: AccessBody):
    rec = CONTRACTS.get(cid)
    if not rec: raise HTTPException(404, "合同不存在")
    if body.access_level not in {"owner", "internal_read", "restricted"}: raise HTTPException(400, "无效权限")
    rec["access_level"] = body.access_level; STORE.save_contract(rec); STORE.audit("access_updated", cid, {"level": body.access_level})
    return {"ok": True, "access_level": body.access_level}


@app.get("/api/audit-log")
def audit_log(contract_id: Optional[str] = None):
    return {"events": STORE.audits(contract_id)}


# ---------- API：CostLens ----------

class CostBody(BaseModel):
    contract_ids: List[str]
    # 多交易公告可按合同 ID 选择具体 transaction_id；未选择时保持阻断而不拼接。
    transaction_ids: Optional[Dict[str, str]] = None


@app.post("/api/analyze/costlens")
def api_costlens(body: CostBody):
    out = []
    for cid in body.contract_ids:
        rec = CONTRACTS.get(cid)
        if not rec:
            continue
        selected = (body.transaction_ids or {}).get(cid)
        parsed = rec["parsed"]
        if selected:
            parsed = costlens.project_transaction(parsed, selected)
            if parsed is None:
                raise HTTPException(400, "合同 %s 未找到所选交易" % rec["name"])
        result = costlens.analyze_cost(parsed)
        out.append({"id": cid, "name": rec["name"], "result": result})
    if not out:
        raise HTTPException(400, "请选择至少一份有效合同")
    return {"analyses": out}


# ---------- API：ScenarioCompare ----------

class ScenarioBody(BaseModel):
    contract_ids: List[str]
    scenario: str = "hold"           # hold / prepay / rate_up / stress / overdue
    month: int = 12
    delta_pp: float = 1.0
    months_late: int = 3
    compare_all: bool = False
    transaction_ids: Optional[Dict[str, str]] = None


@app.post("/api/analyze/scenario")
def api_scenario(body: ScenarioBody):
    if body.scenario not in {"hold", "prepay", "rate_up", "stress", "overdue"}:
        raise HTTPException(400, "未知情景类型，请选择持有到期、提前还款、利率上升、现金流紧张或逾期")
    if body.month < 1 or body.delta_pp < 0 or body.months_late < 0:
        raise HTTPException(400, "情景参数必须为有效的非负数")
    contracts = []
    for cid in body.contract_ids:
        rec = CONTRACTS.get(cid)
        if rec:
            selected = (body.transaction_ids or {}).get(cid)
            parsed = costlens.project_transaction(rec["parsed"], selected) if selected else rec["parsed"]
            if selected and parsed is None:
                raise HTTPException(400, "合同 %s 未找到所选交易" % rec["name"])
            contracts.append({"id": cid, "name": rec["name"], "parsed": parsed})
    if len(contracts) < 1:
        raise HTTPException(400, "请选择合同")
    blocked = []
    for item in contracts:
        gaps = costlens.missing_detail(item["parsed"])
        if gaps:
            blocked.append({"id": item["id"], "name": item["name"], "missing_detail": gaps})
    if blocked:
        return {"mode": "blocked", "reason": "当前信息不足，暂无法进行可靠的情景模拟。",
                "contracts": blocked}
    params = {"month": body.month, "delta_pp": body.delta_pp, "months_late": body.months_late}
    if body.compare_all and len(contracts) >= 2:
        result = scenario.flip_analysis(contracts, params)
        return {"mode": "flip", **result}
    results = scenario.compare_contracts(contracts, body.scenario, params)
    return {"mode": "single", "scenario": body.scenario, "results": results}


# ---------- API：CovenantGuard ----------

class CovenantFin(BaseModel):
    label: str
    total_liabilities: Optional[float] = None
    total_assets: Optional[float] = None
    current_assets: Optional[float] = None
    current_liabilities: Optional[float] = None
    quick_assets: Optional[float] = None


class CovenantBody(BaseModel):
    contract_ids: List[str]
    financials: Optional[List[CovenantFin]] = None
    source_name: str = "用户录入"


@app.post("/api/analyze/covenant")
def api_covenant(body: CovenantBody):
    if not body.financials:
        raise HTTPException(400, "未提供财务数据；CovenantGuard 不会自动使用演示数据")
    fins = [f.dict() for f in body.financials]
    out = []
    for cid in body.contract_ids:
        rec = CONTRACTS.get(cid)
        if not rec:
            continue
        covenants = rec["parsed"].get("covenants", [])
        if not covenants:
            continue
        results = covenant.monitor(covenants, fins)
        out.append({"id": cid, "name": rec["name"], "results": results,
                    "summary": covenant.summary(results),
                    "quarters": [f.get("label") for f in fins]})
        if SHARE_MODE:
            # 公网访客可以放心试算，但录入的财务数据不落库，
            # 避免访客输入污染本地的财务历史与分析快照。
            out[-1]["persisted"] = False
        else:
            STORE.save_financials(cid, body.source_name, fins, out[-1])
            STORE.audit("covenant_recomputed", cid, {"source": body.source_name, "periods": len(fins)})
    if not out:
        raise HTTPException(400, "所选合同未识别到财务契约条款（如资产负债率上限），请换用含契约条款的合同")
    return {"monitors": out}


FIN_HEADER_MAP = {
    "期间": "label", "季度": "label", "日期": "label", "period": "label", "quarter": "label",
    "总负债": "total_liabilities", "负债合计": "total_liabilities", "total liabilities": "total_liabilities",
    "总资产": "total_assets", "资产总计": "total_assets", "total assets": "total_assets",
    "流动资产": "current_assets", "流动资产合计": "current_assets", "current assets": "current_assets",
    "流动负债": "current_liabilities", "流动负债合计": "current_liabilities", "current liabilities": "current_liabilities",
    "速动资产": "quick_assets", "quick assets": "quick_assets", "quick_assets": "quick_assets",
}


def _map_fin_rows(headers, rows):
    mapped = []
    keys = [FIN_HEADER_MAP.get(str(h).strip().lower(), FIN_HEADER_MAP.get(str(h).strip())) for h in headers]
    for row in rows:
        item = {}
        for key, value in zip(keys, row):
            if not key: continue
            if key == "label": item[key] = str(value or "")
            else:
                try: item[key] = float(str(value).replace(",", "")) if value not in (None, "") else None
                except ValueError: item[key] = None
        if item.get("label"): mapped.append(item)
    return mapped


def _normalize_financial_rows(rows):
    """统一 JSON/CSV/XLSX 财务数据类型，无法转换的科目保留为空以便标记无法判断。"""
    if not isinstance(rows, list):
        return rows
    numeric = {"total_liabilities", "total_assets", "current_assets", "current_liabilities", "quick_assets"}
    normalized = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        item = {"label": str(row.get("label") or row.get("期间") or row.get("日期") or "")}
        for key in numeric:
            value = row.get(key)
            if value in (None, ""):
                item[key] = None
                continue
            try:
                item[key] = float(str(value).replace(",", ""))
            except (TypeError, ValueError):
                item[key] = None
        if item["label"]:
            normalized.append(item)
    return normalized


@app.post("/api/contracts/{cid}/financials/upload")
async def upload_financials(cid: str, file: UploadFile = File(...)):
    rec = CONTRACTS.get(cid)
    if not rec: raise HTTPException(404, "合同不存在")
    data = await file.read(); ext = os.path.splitext(file.filename or "")[1].lower()
    if not data: raise HTTPException(400, "财务数据上传失败：文件为空")
    if len(data) > 5 * 1024 * 1024: raise HTTPException(413, "财务数据上传失败：文件超过 5MB 限制")
    try:
        if ext == ".json":
            raw = json.loads(data.decode("utf-8")); fins = raw.get("financials", raw) if isinstance(raw, dict) else raw
            fins = _normalize_financial_rows(fins)
        elif ext == ".csv":
            text = data.decode("utf-8-sig"); table = list(csv.reader(io.StringIO(text))); fins = _map_fin_rows(table[0], table[1:])
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            table = list(ws.iter_rows(values_only=True)); fins = _map_fin_rows(table[0], table[1:])
        else: raise ValueError("仅支持 CSV、JSON、XLSX")
    except Exception as exc:
        raise HTTPException(400, "财务数据解析失败：%s" % exc)
    if not isinstance(fins, list) or not fins: raise HTTPException(400, "未识别到有效财务数据行")
    covenants = rec["parsed"].get("covenants") or []
    results = covenant.monitor(covenants, fins) if covenants else []
    payload = {"id": cid, "name": rec["name"], "results": results, "summary": covenant.summary(results),
               "financials": fins, "source_name": file.filename, "mapped_headers": sorted({k for f in fins for k in f})}
    STORE.save_financials(cid, file.filename, fins, payload); STORE.audit("financial_file_uploaded", cid, {"name": file.filename, "rows": len(fins)})
    return payload


@app.get("/api/contracts/{cid}/financials/history")
def financial_history(cid: str):
    if cid not in CONTRACTS: raise HTTPException(404, "合同不存在")
    return {"history": STORE.financial_history(cid)}


# ---------- API：ContractDiff ----------

class DiffBody(BaseModel):
    old_id: str
    new_id: str
    remaining_principal: Optional[float] = None


@app.post("/api/analyze/diff")
def api_diff(body: DiffBody):
    if body.old_id == body.new_id:
        raise HTTPException(400, "请选择两份不同的合同版本")
    old = CONTRACTS.get(body.old_id)
    new = CONTRACTS.get(body.new_id)
    if not old or not new:
        raise HTTPException(400, "请选择两份合同（旧版与新版）")
    result = diffengine.diff_contracts(old["parsed"], new["parsed"],
                                       remaining_principal=body.remaining_principal)
    return {
        "old": {"id": old["id"], "name": old["name"]},
        "new": {"id": new["id"], "name": new["name"]},
        **result,
    }


# ---------- API：四引擎综合评估 ----------

class ComprehensiveBody(BaseModel):
    contract_ids: List[str]
    prepay_month: int = 12
    financials: Optional[List[CovenantFin]] = None


@app.post("/api/analyze/comprehensive")
def api_comprehensive(body: ComprehensiveBody):
    records = [CONTRACTS[cid] for cid in body.contract_ids if cid in CONTRACTS]
    if len(records) < 2:
        raise HTTPException(400, "一键综合评估至少需要两份有效合同")
    fins = [f.dict() for f in body.financials] if body.financials else []
    try:
        return assessment.analyze_bundle(records, fins, max(1, body.prepay_month))
    except ValueError as exc:
        raise HTTPException(400, str(exc))


# ---------- API：任务驱动 Agent ----------

# 可用分析工具白名单：模型输出必须落在这个集合内，否则丢弃
AGENT_TOOLS = ("risk", "costlens", "scenario", "diff", "covenant")
# 意图分类单独设超时，不与字段补强的 25 秒共用；超时即回退，不让用户干等
AGENT_LLM_TIMEOUT = 6

# 关键词回退表：公网只读模式（deterministic_only=True）与无 Key 时的唯一可用路径
AGENT_INTENT_KEYWORDS = {
    "risk": ("审查", "风险", "隐患", "坑", "违约", "复核", "谈判", "不利",
             "有没有问题", "帮我看看", "审一下"),
    "costlens": ("成本", "融资", "利息", "贵", "便宜", "划算", "值不值", "利率",
                 "负担", "费用高"),
    "scenario": ("比较", "方案", "对比", "哪个好", "选哪个", "优劣", "更优",
                 "二选一", "更合适", "哪份"),
    "diff": ("版本", "变化", "新旧", "改了什么", "区别", "差异", "变更", "两版", "改了哪"),
    "covenant": ("契约", "财务", "监控", "承诺", "指标", "阈值", "触发", "达标", "约束"),
}

AGENT_INTENT_PROMPT = """你是金融合同分析助手。根据用户任务，判断需要调用哪些分析工具。
可用工具：
- risk：风险识别（违约、罚息、担保、合规等）
- costlens：融资成本测算（利率、费用、真实年化）
- scenario：多方案比较（仅当选中 2 份及以上合同）
- diff：版本差异（仅当恰好选中 2 份合同）
- covenant：财务契约监控（仅在提供财务数据时）

用户任务：{task}
选中合同数：{n}
是否提供财务数据：{has_fin}

只返回 JSON 数组，元素为工具名，例如 ["costlens"]。无法判断时返回 []。不要任何解释。"""


def _agent_keywords(task):
    """关键词命中：返回工具名列表（可能含多个）。"""
    return [tool for tool, words in AGENT_INTENT_KEYWORDS.items()
            if any(word in task for word in words)]


def _agent_normalize(tools, n_contracts, has_financials):
    """白名单过滤 + 上下文约束。模型输出与关键词结果都必须过这里。"""
    valid = []
    for tool in tools:
        if tool in AGENT_TOOLS and tool not in valid:
            valid.append(tool)
    if n_contracts < 2:
        # 单份合同谈不上方案比较与版本对比
        valid = [t for t in valid if t not in ("scenario", "diff")]
    elif n_contracts != 2:
        # 三份及以上不能当版本对比，避免把不同方案误判为新旧版本
        valid = [t for t in valid if t != "diff"]
    if not has_financials:
        valid = [t for t in valid if t != "covenant"]
    return valid


def _agent_llm_tools(task, n_contracts, has_financials):
    """让模型出计划。任何失败（超时/无输出/非 JSON）返回 None，交给调用方回退。"""
    try:
        prompt = AGENT_INTENT_PROMPT.format(
            task=task, n=n_contracts, has_fin="是" if has_financials else "否")
        text = ai_client.complete([{"role": "user", "content": prompt}],
                                  max_tokens=64, temperature=0.0, timeout=AGENT_LLM_TIMEOUT)
        if not text:
            return None
        matched = re.search(r"\[[^\]]*\]", text, re.S)
        if not matched:
            return None
        arr = json.loads(matched.group(0))
        if not isinstance(arr, list):
            return None
        return [str(item).strip().lower() for item in arr if isinstance(item, str)]
    except Exception:
        return None


def _agent_plan(task, n_contracts, has_financials):
    """返回 (工具列表, 计划来源)。来源取值：llm / keyword / fallback。"""
    if SECURITY["deterministic_only"] or not ai_client.is_configured():
        return _agent_normalize(_agent_keywords(task), n_contracts, has_financials), "keyword"
    raw = _agent_llm_tools(task, n_contracts, has_financials)
    if raw:
        plan = _agent_normalize(raw, n_contracts, has_financials)
        if plan:
            return plan, "llm"
        # 模型给出的工具全被白名单或上下文约束剔除：视为无效计划，
        # 回退关键词而不是标着 llm 却什么都不做
    return _agent_normalize(_agent_keywords(task), n_contracts, has_financials), "fallback"


RISK_CONTEXT_MARK = "【风险审查】"


def _build_ai_context(results):
    """从本次执行的 risks 里提炼可读结论，供聊天框与 Agent 页共用。

    返回 (文本, 结构化列表)。文本以 RISK_CONTEXT_MARK 开头，
    内置引擎据此识别"这是风险段"，按四要素回答而不是逐行复述。

    证据字段用 source_text：riskengine 里的 evidence 恒为 None，
    source_text 才是逐字摘录的原文。
    """
    order = {"高": 0, "中": 1, "低": 2}
    items = []
    for rec in results.get("reviews") or []:
        for risk in sorted(rec.get("risks") or [], key=lambda x: order.get(x.get("level"), 9))[:3]:
            items.append({
                "contract": rec.get("name"),
                "title": risk.get("title"),
                "level": risk.get("level"),
                "why": risk.get("explanation"),
                "evidence": risk.get("source_text"),
                "clause": risk.get("clause"),
                "page": risk.get("page"),
                "advice": risk.get("suggestion"),
                "legal_review": bool(risk.get("needs_legal_review")),
            })
    if not items:
        return "", []
    lines = [RISK_CONTEXT_MARK]
    for it in items:
        lines.append("【%s】%s（%s风险）" % (it["contract"] or "合同", it["title"] or "风险项", it["level"] or "未分级"))
        lines.append("为什么：%s" % (it["why"] or "—"))
        where = "、".join(x for x in (it["clause"], ("第 %s 页" % it["page"]) if it["page"] else None) if x)
        lines.append("依据：%s%s" % (it["evidence"] or "—", ("（%s）" % where) if where else ""))
        lines.append("建议：%s" % (it["advice"] or "—"))
        if it["legal_review"]:
            lines.append("提示：该项建议法务复核")
        lines.append("")
    return "\n".join(lines).strip(), items


class AgentTaskBody(BaseModel):
    task: str
    contract_ids: List[str]
    financials: Optional[List[CovenantFin]] = None


@app.post("/api/agent/tasks")
def run_agent_task(body: AgentTaskBody):
    records = [CONTRACTS[cid] for cid in body.contract_ids if cid in CONTRACTS]
    if not records: raise HTTPException(400, "请先选择合同")
    task = body.task.strip()
    tool_trace, results, review = [], {}, []

    def trace(tool, status, why, kind="deterministic"):
        tool_trace.append({"tool": tool, "status": status, "reason": why, "result_type": kind})

    tools, plan_source = _agent_plan(task, len(records), bool(body.financials))
    trace("Plan", "called",
          {"llm": "由大模型识别任务意图并生成执行计划",
           "keyword": "未启用大模型，按关键词规则生成执行计划",
           "fallback": "大模型不可用或输出不合法，已回退关键词规则"}[plan_source],
          kind="ai_judgment" if plan_source == "llm" else "deterministic")

    if "risk" in tools:
        agents = []
        for rec in records:
            out = orchestrator.orchestrate(rec["parsed"]); agents.append({"id": rec["id"], "name": rec["name"], **out})
            review.extend(out.get("review", []))
        results["reviews"] = agents
        trace("Risk Engine", "called", "任务需要识别风险、证据与人工复核项")
    else:
        trace("Risk Engine", "skipped", "当前目标未要求完整风险审查")

    if "costlens" in tools:
        analyses = []
        for rec in records:
            r = costlens.analyze_cost(rec["parsed"])
            if not r.get("hidden"): analyses.append({"id": rec["id"], "name": rec["name"], "result": r})
        if analyses:
            results["costlens"] = analyses; trace("Cost Lens", "called", "所选合同属于融资类且可计算要素齐备")
        else: trace("Cost Lens", "skipped", "合同类型不产生融资现金流或关键要素不足")
    else: trace("Cost Lens", "skipped", "用户目标不涉及融资成本")

    if "scenario" in tools:
        wrapped = [{"id": r["id"], "name": r["name"], "parsed": r["parsed"]} for r in records]
        try:
            scenario_result = scenario.flip_analysis(wrapped, {"month": 12})
            if scenario_result.get("blocked"):
                trace("Scenario Compare", "degraded", "部分合同缺少可确认的实际现金流要素，已阻止排名比较")
            else:
                results["scenario"] = scenario_result
                trace("Scenario Compare", "called", "检测到多份融资合同，需要统一情景比较")
        except Exception:
            trace("Scenario Compare", "skipped", "部分合同缺少可计算要素")
    else: trace("Scenario Compare", "skipped", "需要至少两份可测算合同且任务为方案比较")

    if "diff" in tools:
        results["diff"] = diffengine.diff_contracts(records[0]["parsed"], records[1]["parsed"])
        trace("Contract Diff", "called", "用户明确要求检查两份合同的版本变化")
    else: trace("Contract Diff", "skipped", "仅在两份文本且目标为版本变化时调用，避免把不同方案误当版本")

    if "covenant" in tools:
        fins = [f.dict() for f in body.financials]
        monitored = []
        for rec in records:
            if rec["parsed"].get("covenants"):
                rs = covenant.monitor(rec["parsed"]["covenants"], fins)
                monitored.append({"id": rec["id"], "name": rec["name"], "results": rs, "summary": covenant.summary(rs)})
        results["covenant"] = monitored; trace("Covenant Guard", "called" if monitored else "skipped", "使用用户提供的财务数据计算合同契约指标")
    else: trace("Covenant Guard", "skipped", "未提供财务数据或任务不涉及契约监控")

    ai_used = ai_client.is_configured() and not SECURITY["deterministic_only"]
    trace("External LLM", "available" if ai_used else "skipped", "仅用于语义解释；关键金融计算均由确定性代码完成", "ai_judgment")
    ai_context, ai_summary = _build_ai_context(results)
    STORE.audit("agent_task_run", records[0]["id"], {"task": task, "contracts": [r["id"] for r in records], "tools": tool_trace, "plan_source": plan_source})
    return {"task": task, "contracts": [{"id": r["id"], "name": r["name"]} for r in records],
            "tool_trace": tool_trace, "results": results, "human_review": review,
            "plan_source": plan_source,
            "ai_context": ai_context, "ai_summary": ai_summary,
            "external_llm": {"used": ai_used, "scope": SECURITY["external_scope"] if ai_used else "none"}}


class SecurityBody(BaseModel):
    deterministic_only: bool
    external_scope: str = "masked_evidence_only"


@app.get("/api/security")
def security_status():
    return {**SECURITY, "local_processing": True, "external_model_configured": ai_client.is_configured(),
            "external_data": "脱敏后的必要证据片段" if not SECURITY["deterministic_only"] else "不发送"}


@app.post("/api/security")
def update_security(body: SecurityBody):
    SECURITY.update({"deterministic_only": body.deterministic_only, "external_scope": body.external_scope})
    STORE.audit("security_setting_updated", None, SECURITY)
    return {"ok": True, **SECURITY}


# ---------- API：AI ----------

# 模型标识通常形如 gpt-4o-mini / deepseek-chat / glm-4-flash，限定字符集即可，不限定具体供应商
_MODEL_NAME_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:\-/]{0,63}$")


def _validate_ai_config(model, proxy):
    """校验 AI 配置的可写字段。

    原实现把请求体直接写进 config.json：模型名曾被模糊测试改写成随机字符串，
    代理则可指向任意地址，影响后续全部外部请求。这里只做格式约束，
    不锁定具体模型，避免堵死自定义部署。
    """
    if model and not _MODEL_NAME_RE.match(model.strip()):
        raise HTTPException(400, "模型名称不合法：仅允许字母、数字与 . _ : - /，最长 64 字符")
    if proxy:
        parsed = urlparse(proxy.strip())
        if parsed.scheme not in ("http", "https") or not parsed.netloc:
            raise HTTPException(400, "代理地址不合法：需形如 http://127.0.0.1:7890")


class AIConfigBody(BaseModel):
    provider: Optional[str] = None
    model: Optional[str] = None
    proxy: Optional[str] = None
    test: bool = False


class AIChatBody(BaseModel):
    messages: List[dict]
    context: Optional[str] = None


@app.get("/api/ai/status")
def ai_status():
    return ai_client.status()


@app.post("/api/ai/config")
def ai_config(body: AIConfigBody):
    _validate_ai_config(body.model, body.proxy)
    if body.provider:
        ai_client.set_provider(body.provider)
    if body.model:
        ai_client.CONFIG["model"] = body.model.strip()
    if body.proxy is not None:
        ai_client.CONFIG["proxy"] = body.proxy.strip()
    ai_client.save_config(CONFIG_PATH)
    st = ai_client.status()
    test_result = None
    if body.test and st["configured"]:
        test_result = ai_client.test_connection()
    return {"ok": True, "status": st, "test": test_result}


@app.post("/api/ai/chat")
def ai_chat(body: AIChatBody):
    if not body.messages or len(body.messages) > 30:
        raise HTTPException(400, "消息数量必须为 1–30 条")
    if sum(len(str(m.get("content", ""))) for m in body.messages) > 40_000:
        raise HTTPException(400, "对话内容过长，请精简后重试")
    def gen():
        try:
            first = True
            stream = (ai_client._builtin_engine_stream(body.messages, body.context)
                      if SECURITY["deterministic_only"] else ai_client.chat_stream(body.messages, file_parser.mask_sensitive(body.context or "")))
            for chunk in stream:
                if first:
                    yield "data: " + json.dumps({"type": "meta", "engine": "builtin" if SECURITY["deterministic_only"] else ("llm" if ai_client.is_configured() else "builtin")},
                                                 ensure_ascii=False) + "\n\n"
                    first = False
                yield "data: " + json.dumps({"type": "delta", "content": chunk}, ensure_ascii=False) + "\n\n"
            yield "data: " + json.dumps({"type": "done"}, ensure_ascii=False) + "\n\n"
        except Exception as e:
            yield "data: " + json.dumps({"type": "error", "message": str(e)[:200]}, ensure_ascii=False) + "\n\n"

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.post("/api/ai/chat-once")
def ai_chat_once(body: AIChatBody):
    """Non-streaming fallback for Cloudflare Quick Tunnel, which does not support SSE."""
    if not body.messages or len(body.messages) > 30:
        raise HTTPException(400, "消息数量必须为 1–30 条")
    if sum(len(str(m.get("content", ""))) for m in body.messages) > 40_000:
        raise HTTPException(400, "对话内容过长，请精简后重试")
    try:
        stream = (ai_client._builtin_engine_stream(body.messages, body.context)
                  if SECURITY["deterministic_only"] else
                  ai_client.chat_stream(body.messages, file_parser.mask_sensitive(body.context or "")))
        return {"content": "".join(stream), "engine": "builtin" if SECURITY["deterministic_only"]
                else ("llm" if ai_client.is_configured() else "builtin")}
    except Exception as exc:
        raise HTTPException(502, "AI 服务暂时不可用：%s" % str(exc)[:160])


@app.get("/api/health")
def health():
    """用于演示环境快速检查：确定性引擎始终可用，AI 可独立降级。"""
    st = ai_client.status()
    db = STORE.health()
    return {"ok": db["ok"], "mode": "share" if SHARE_MODE else ("demo" if DEMO_MODE else "local"),
            "demo_mode": DEMO_MODE, "share_mode": SHARE_MODE, "contracts": len(CONTRACTS),
            "readonly": SHARE_MODE, "write_token_configured": bool(SHARE_TOKEN),
            "database": db, "ai": {k: st[k] for k in ("configured", "provider", "label", "model")}}


# ---------- 静态资源 ----------

app.mount("/", StaticFiles(directory=os.path.join(BASE_DIR, "static"), html=True), name="static")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host=os.getenv("CONTRACTGUARD_HOST", "127.0.0.1"),
                port=int(os.getenv("CONTRACTGUARD_PORT", "8899")))
